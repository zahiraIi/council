"""Plain text out of the source formats, as a list of (page_label, text) units.
page_label is a PDF page number, a slide number, or a markdown section heading."""
import re


def pdf_pages(path):
    from pypdf import PdfReader
    out = []
    for i, page in enumerate(PdfReader(path).pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            out.append((i, text))
    return out


def docx_pages(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))   # merged cells repeat; dedupe
    text = "\n".join(parts)
    return [(1, text)] if text.strip() else []


def pptx_pages(path):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = "\n".join(p for p in parts if p.strip())
        if text.strip():
            out.append((i, text))
    return out


def md_sections(text, level=2):
    """Split markdown on headings of `level` or higher. Returns [(heading, body)];
    text before the first heading is labelled 'intro'."""
    pat = re.compile(r"^(#{1,%d})\s+(.*)$" % level, re.M)
    out, last, head = [], 0, "intro"
    for m in pat.finditer(text):
        body = text[last:m.start()].strip()
        if body:
            out.append((head, body))
        head, last = m.group(2).strip(), m.start()
    body = text[last:].strip()
    if body:
        out.append((head, body))
    return out


def md_pages(path):
    with open(path, encoding="utf-8") as f:
        return md_sections(f.read())


def pages(path):
    ext = path.rsplit(".", 1)[-1].lower()
    fn = {"pdf": pdf_pages, "docx": docx_pages, "pptx": pptx_pages, "md": md_pages}.get(ext)
    return fn(path) if fn else []


def chunk_pages(units, size=260, overlap=40):
    """Word-window chunks across page units, so a chunk can span a page break.
    Returns [{"text", "first", "last"}] where first/last are page labels.
    ~260 words ≈ 350 tokens; `overlap` words repeat between neighbours."""
    words = []
    for label, text in units:
        words += [(w, label) for w in text.split()]
    if not words:
        return []
    step = max(1, size - overlap)
    out = []
    for start in range(0, len(words), step):
        win = words[start:start + size]
        out.append({"text": " ".join(w for w, _ in win), "first": win[0][1], "last": win[-1][1]})
        if start + size >= len(words):
            break
    return out
